import os
import re
import shutil
import logging
import asyncio
from datetime import datetime
from pathlib import Path

from dotenv import load_dotenv
from openpyxl import Workbook, load_workbook
from telegram import Update
from telegram.error import Forbidden, RetryAfter, TelegramError
from telegram.ext import (
    Application,
    CommandHandler,
    ConversationHandler,
    MessageHandler,
    ContextTypes,
    filters,
)

# =========================================================
# GETE REGISTRATION BOT
# =========================================================

BASE_DIR = Path(__file__).resolve().parent
load_dotenv(BASE_DIR / ".env")

BOT_TOKEN = os.getenv("BOT_TOKEN", "").strip()

# Supports one administrator (ADMIN_ID) or several explicitly configured
# administrators (ADMIN_IDS=123,456). ADMIN_IDS takes precedence.
_admin_value = os.getenv("ADMIN_IDS", "").strip()
if not _admin_value:
    _admin_value = os.getenv("ADMIN_ID", "").strip()

if not BOT_TOKEN or not _admin_value:
    raise RuntimeError("Missing BOT_TOKEN or ADMIN_ID/ADMIN_IDS. Check your .env file.")

try:
    ADMIN_IDS = {int(item.strip()) for item in _admin_value.split(",") if item.strip()}
    if not ADMIN_IDS:
        raise ValueError
except ValueError as exc:
    raise RuntimeError("ADMIN_ID/ADMIN_IDS must contain numeric Telegram user IDs.") from exc

EXCEL_FILE = BASE_DIR / "students.xlsx"

BASE_HEADERS = [
    "Student ID",
    "Full Name",
    "Phone Number",
    "Current Address",
    "Registration Date",
]
TELEGRAM_HEADERS = ["Telegram User ID", "Telegram Chat ID"]
HEADERS = BASE_HEADERS + TELEGRAM_HEADERS

NAME, PHONE, ADDRESS, EDIT_FIELD, EDIT_VALUE, BROADCAST_MESSAGE, BROADCAST_CONFIRM, REGISTRATION_CONFIRM = range(8)

logging.basicConfig(
    format="%(asctime)s - %(name)s - %(levelname)s - %(message)s",
    level=logging.INFO,
)
logging.getLogger("httpx").setLevel(logging.WARNING)
logger = logging.getLogger(__name__)


# =========================================================
# EXCEL HELPERS
# =========================================================

def _header_map(sheet):
    """Map header labels to Excel column numbers (1-based)."""
    result = {}
    for column in range(1, sheet.max_column + 1):
        value = sheet.cell(1, column).value
        if value is not None:
            result[str(value).strip()] = column
    return result


def ensure_excel():
    """Create workbook or safely add missing columns without removing records."""
    if not EXCEL_FILE.exists():
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Students"
        sheet.append(HEADERS)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        workbook.save(EXCEL_FILE)
        workbook.close()
        logger.info("Created Excel file: %s", EXCEL_FILE)
        return

    workbook = load_workbook(EXCEL_FILE)
    if "Students" in workbook.sheetnames:
        sheet = workbook["Students"]
    else:
        sheet = workbook.active
        sheet.title = "Students"

    if sheet.cell(1, 1).value is None:
        for column, header in enumerate(BASE_HEADERS, start=1):
            sheet.cell(1, column, header)
    elif str(sheet.cell(1, 1).value).strip() != "Student ID":
        workbook.close()
        raise RuntimeError(
            "The existing Excel workbook has unexpected columns. "
            "Back it up and check its headers before starting the bot."
        )

    mapping = _header_map(sheet)
    # Keep the original five-column layout intact and append new fields
    # only when they are missing. Existing data is not deleted.
    for header in BASE_HEADERS:
        if header not in mapping:
            workbook.close()
            raise RuntimeError(
                f"Required Excel header '{header}' is missing. "
                "Back up students.xlsx and check its first five columns."
            )

    for header in TELEGRAM_HEADERS:
        if header not in mapping:
            column = sheet.max_column + 1
            sheet.cell(1, column, header)
            mapping[header] = column

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    workbook.save(EXCEL_FILE)
    workbook.close()


def _as_int(value):
    try:
        return int(value)
    except (TypeError, ValueError):
        try:
            return int(float(value))
        except (TypeError, ValueError):
            return None


def read_students():
    """Read student rows using header names, preserving existing records."""
    ensure_excel()
    workbook = load_workbook(EXCEL_FILE, read_only=True, data_only=True)
    sheet = workbook["Students"]
    mapping = _header_map(sheet)
    students = []

    for row_number in range(2, sheet.max_row + 1):
        student_id = _as_int(sheet.cell(row_number, mapping["Student ID"]).value)
        if student_id is None:
            continue

        def cell(header):
            col = mapping.get(header)
            return sheet.cell(row_number, col).value if col else None

        students.append({
            "id": student_id,
            "name": str(cell("Full Name") or ""),
            "phone": str(cell("Phone Number") or ""),
            "address": str(cell("Current Address") or ""),
            "date": str(cell("Registration Date") or ""),
            "telegram_user_id": _as_int(cell("Telegram User ID")),
            "chat_id": _as_int(cell("Telegram Chat ID")),
        })

    workbook.close()
    return students


def find_student_by_telegram_id(telegram_user_id):
    return next(
        (s for s in read_students() if s["telegram_user_id"] == telegram_user_id),
        None,
    )


def add_student(name, phone, address, telegram_user_id, chat_id):
    """Append a student; prevent a second row for the same Telegram user."""
    ensure_excel()
    workbook = load_workbook(EXCEL_FILE)
    sheet = workbook["Students"]
    mapping = _header_map(sheet)

    # Avoid duplicate registrations by Telegram user ID.
    for row_number in range(2, sheet.max_row + 1):
        saved_id = _as_int(
            sheet.cell(row_number, mapping["Telegram User ID"]).value
        )
        if saved_id == telegram_user_id:
            existing_student_id = _as_int(
                sheet.cell(row_number, mapping["Student ID"]).value
            )
            workbook.close()
            return existing_student_id, None, True

    existing_ids = [
        student_id for student_id in (
            _as_int(sheet.cell(row, mapping["Student ID"]).value)
            for row in range(2, sheet.max_row + 1)
        ) if student_id is not None
    ]
    student_id = max(existing_ids, default=0) + 1
    registered_at = datetime.now().astimezone().strftime("%Y-%m-%d %H:%M")

    values = {
        "Student ID": student_id,
        "Full Name": name,
        "Phone Number": phone,
        "Current Address": address,
        "Registration Date": registered_at,
        "Telegram User ID": telegram_user_id,
        "Telegram Chat ID": chat_id,
    }
    row_number = sheet.max_row + 1
    for header, value in values.items():
        sheet.cell(row_number, mapping[header], value)

    sheet.auto_filter.ref = sheet.dimensions
    workbook.save(EXCEL_FILE)
    workbook.close()
    return student_id, registered_at, False


def get_student(student_id):
    return next((s for s in read_students() if s["id"] == student_id), None)


def edit_student(student_id, field, value):
    columns = {
        "name": "Full Name",
        "phone": "Phone Number",
        "address": "Current Address",
    }
    if field not in columns:
        return False

    workbook = load_workbook(EXCEL_FILE)
    sheet = workbook["Students"]
    mapping = _header_map(sheet)

    for row_number in range(2, sheet.max_row + 1):
        current_id = _as_int(sheet.cell(row_number, mapping["Student ID"]).value)
        if current_id == student_id:
            sheet.cell(row_number, mapping[columns[field]], value)
            workbook.save(EXCEL_FILE)
            workbook.close()
            return True

    workbook.close()
    return False


def remove_student(student_id):
    workbook = load_workbook(EXCEL_FILE)
    sheet = workbook["Students"]
    mapping = _header_map(sheet)

    for row_number in range(2, sheet.max_row + 1):
        current_id = _as_int(sheet.cell(row_number, mapping["Student ID"]).value)
        if current_id == student_id:
            sheet.delete_rows(row_number, 1)
            sheet.auto_filter.ref = sheet.dimensions
            workbook.save(EXCEL_FILE)
            workbook.close()
            return True

    workbook.close()
    return False


# =========================================================
# COMMON HELPERS
# =========================================================

def valid_phone(phone):
    phone = phone.strip()
    if not re.fullmatch(r"\+?[0-9][0-9\s-]{5,18}[0-9]", phone):
        return False
    digits = re.sub(r"\D", "", phone)
    return 7 <= len(digits) <= 15


def clean_text(text):
    return " ".join(text.strip().split())


def is_admin(update):
    user = update.effective_user
    return user is not None and user.id in ADMIN_IDS


async def admin_only(update):
    if not is_admin(update):
        await update.effective_message.reply_text(
            "Access denied. This command is for authorized administrators only."
        )
        return False
    if update.effective_chat.type != "private":
        await update.effective_message.reply_text(
            "Please use administration commands in a private chat."
        )
        return False
    return True


async def send_long_message(message, text, limit=3500):
    while len(text) > limit:
        split_at = text.rfind("\n", 0, limit)
        if split_at <= 0:
            split_at = limit
        await message.reply_text(text[:split_at])
        text = text[split_at:].lstrip("\n")
    if text:
        await message.reply_text(text)


# =========================================================
# START / HELP / ADMIN
# =========================================================

async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    message = update.effective_message
    user = update.effective_user
    chat = update.effective_chat

    if chat.type != "private":
        await message.reply_text("Please open this bot in a private chat to register.")
        return ConversationHandler.END

    # Do not create duplicate records for an already registered Telegram user.
    existing = find_student_by_telegram_id(user.id)
    if existing:
        await message.reply_text(
            f"Welcome back, {existing['name']}! 🎓\n\n"
            f"You are already registered.\nStudent ID: {existing['id']}\n\n"
            "If your details need to change, contact the administrator."
        )
        return ConversationHandler.END

    context.user_data.pop("registration", None)
    await message.reply_text(
        "Welcome to Gete Registration Bot! 🎓\n\n"
        "Let's register you now.\n\n"
        "Step 1 of 3: Please enter your full name.\n\n"
        "Use /cancel at any time to stop."
    )
    return NAME


async def help_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.effective_message.reply_text(
        "GETE REGISTRATION BOT\n\n"
        "STUDENT\n"
        "/start - Begin registration automatically\n"
        "/register - Start registration\n"
        "/cancel - Cancel the current operation\n"
        "/help - Show this help\n\n"
        "ADMIN\n"
        "/admin - Show admin controls\n"
        "/students - List registered students\n"
        "/search WORD - Search student records\n"
        "/edit ID - Edit a student's details\n"
        "/delete ID - Delete a student record\n"
        "/export - Download the Excel file\n"
        "/broadcast - Send an announcement to registered students\n\n"
        "Admin commands work only for configured admin IDs in private chat."
    )


async def admin_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not await admin_only(update):
        return
    await update.effective_message.reply_text(
        "GETE REGISTRATION — ADMIN CONTROLS\n\n"
        "/students — list student records\n"
        "/search WORD — search records\n"
        "/edit ID — edit student details\n"
        "/delete ID — delete a student\n"
        "/export — download students.xlsx\n"
        "/broadcast — prepare an announcement\n\n"
        "Broadcasts require confirmation before sending."
    )


# =========================================================
# STUDENT REGISTRATION CONVERSATION
# =========================================================

async def register(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if update.effective_chat.type != "private":
        await update.effective_message.reply_text(
            "Please register through a private chat with this bot."
        )
        return ConversationHandler.END

    existing = find_student_by_telegram_id(update.effective_user.id)
    if existing:
        await update.effective_message.reply_text(
            f"You are already registered as {existing['name']} "
            f"(Student ID: {existing['id']})."
        )
        return ConversationHandler.END

    context.user_data.pop("registration", None)
    await update.effective_message.reply_text(
        "STUDENT REGISTRATION\n\nStep 1 of 3\nEnter your full name:"
    )
    return NAME


async def receive_name(update: Update, context: ContextTypes.DEFAULT_TYPE):
    name = clean_text(update.effective_message.text)
    if not 2 <= len(name) <= 100:
        await update.effective_message.reply_text(
            "Please enter a full name between 2 and 100 characters."
        )
        return NAME

    context.user_data["registration"] = {"name": name}
    await update.effective_message.reply_text(
        "Step 2 of 3\nEnter your phone number.\nExample: +251912345678"
    )
    return PHONE


async def receive_phone(update: Update, context: ContextTypes.DEFAULT_TYPE):
    phone = update.effective_message.text.strip()
    if not valid_phone(phone):
        await update.effective_message.reply_text(
            "Invalid phone number. Please try again.\nExample: +251912345678"
        )
        return PHONE

    context.user_data["registration"]["phone"] = phone
    await update.effective_message.reply_text(
        "Step 3 of 3\nEnter your current address (town, city, or neighborhood)."
    )
    return ADDRESS


async def receive_address(update: Update, context: ContextTypes.DEFAULT_TYPE):
    address = clean_text(update.effective_message.text)
    if not 2 <= len(address) <= 200:
        await update.effective_message.reply_text(
            "Please enter an address between 2 and 200 characters."
        )
        return ADDRESS

    data = context.user_data.get("registration")
    if not data:
        await update.effective_message.reply_text(
            "Your registration session expired. Send /start to begin again."
        )
        return ConversationHandler.END

    data["address"] = address
    await update.effective_message.reply_text(
        "Please confirm your registration details:\n\n"
        f"Full name: {data['name']}\n"
        f"Phone number: {data['phone']}\n"
        f"Current address: {address}\n\n"
        "Reply YES to save your registration, or NO to cancel."
    )
    return REGISTRATION_CONFIRM


async def confirm_registration(
    update: Update,
    context: ContextTypes.DEFAULT_TYPE,
):
    answer = update.effective_message.text.strip().casefold()
    if answer in {"no", "n", "cancel"}:
        context.user_data.pop("registration", None)
        await update.effective_message.reply_text(
            "Registration cancelled. Send /start whenever you are ready."
        )
        return ConversationHandler.END

    if answer not in {"yes", "y"}:
        await update.effective_message.reply_text(
            "Please reply YES to save your details or NO to cancel."
        )
        return REGISTRATION_CONFIRM

    data = context.user_data.get("registration")
    if not data or not data.get("address"):
        await update.effective_message.reply_text(
            "Your registration session expired. Send /start to begin again."
        )
        return ConversationHandler.END

    user = update.effective_user
    chat = update.effective_chat
    try:
        student_id, registered_at, duplicate = add_student(
            data["name"], data["phone"], data["address"], user.id, chat.id
        )
    except PermissionError:
        await update.effective_message.reply_text(
            "I cannot write to students.xlsx because it may be open in Excel. "
            "Close the workbook and send /start to try again."
        )
        return ConversationHandler.END
    except Exception:
        logger.exception("Unable to save student registration")
        await update.effective_message.reply_text(
            "Sorry, the registration could not be saved. Please contact the administrator."
        )
        return ConversationHandler.END

    context.user_data.pop("registration", None)

    if duplicate:
        await update.effective_message.reply_text(
            f"You are already registered (Student ID: {student_id})."
        )
        return ConversationHandler.END

    await update.effective_message.reply_text(
        "✅ REGISTRATION SUCCESSFUL!\n\n"
        f"Student ID: {student_id}\n"
        f"Full name: {data['name']}\n"
        f"Phone number: {data['phone']}\n"
        f"Current address: {data['address']}\n"
        f"Registration date: {registered_at}\n\n"
        "Your information has been saved."
    )
    logger.info("Saved student ID %s from Telegram user %s", student_id, user.id)
    return ConversationHandler.END


async def cancel(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data.pop("registration", None)
    context.user_data.pop("edit_student_id", None)
    context.user_data.pop("edit_field", None)
    context.user_data.pop("broadcast_message", None)
    await update.effective_message.reply_text(
        "Operation cancelled. Send /start to register or /help for commands."
    )
    return ConversationHandler.END


# =========================================================
# ADMIN: STUDENT LIST / SEARCH / EDIT / DELETE / EXPORT
# =========================================================

async def students_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not await admin_only(update):
        return
    students = read_students()
    if not students:
        await update.effective_message.reply_text("No student registrations found.")
        return

    await update.effective_message.reply_text(f"Total registered students: {len(students)}")
    lines = [
        f"ID: {s['id']}\nName: {s['name']}\nPhone: {s['phone']}\n"
        f"Address: {s['address']}\nRegistered: {s['date']}\n"
        "-------------------------"
        for s in students
    ]
    await send_long_message(update.effective_message, "\n".join(lines))


async def search_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not await admin_only(update):
        return
    query = " ".join(context.args).strip().casefold()
    if not query:
        await update.effective_message.reply_text("Enter a search term. Example: /search Mekelle")
        return

    matches = [
        s for s in read_students()
        if query in f"{s['id']} {s['name']} {s['phone']} {s['address']} {s['date']}".casefold()
    ]
    if not matches:
        await update.effective_message.reply_text("No matching students found.")
        return

    lines = [
        f"ID: {s['id']}\nName: {s['name']}\nPhone: {s['phone']}\n"
        f"Address: {s['address']}\nRegistered: {s['date']}\n"
        "-------------------------"
        for s in matches
    ]
    await update.effective_message.reply_text(f"Found {len(matches)} matching student(s).")
    await send_long_message(update.effective_message, "\n".join(lines))


async def edit_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not await admin_only(update):
        return ConversationHandler.END
    if not context.args or not context.args[0].isdigit():
        await update.effective_message.reply_text("Usage: /edit STUDENT_ID\nExample: /edit 1")
        return ConversationHandler.END

    student_id = int(context.args[0])
    student = get_student(student_id)
    if not student:
        await update.effective_message.reply_text("Student ID not found.")
        return ConversationHandler.END

    context.user_data["edit_student_id"] = student_id
    await update.effective_message.reply_text(
        f"Editing student {student_id}: {student['name']}\n\n"
        "Which field should change? Send: name, phone, or address.\n"
        "Use /cancel to stop."
    )
    return EDIT_FIELD


async def receive_edit_field(update: Update, context: ContextTypes.DEFAULT_TYPE):
    field = update.effective_message.text.strip().lower()
    if field not in {"name", "phone", "address"}:
        await update.effective_message.reply_text("Please enter name, phone, or address.")
        return EDIT_FIELD
    context.user_data["edit_field"] = field
    await update.effective_message.reply_text(f"Enter the new {field}:")
    return EDIT_VALUE


async def receive_edit_value(update: Update, context: ContextTypes.DEFAULT_TYPE):
    field = context.user_data.get("edit_field")
    value = clean_text(update.effective_message.text)
    student_id = context.user_data.get("edit_student_id")

    if not field or not student_id:
        await update.effective_message.reply_text("Edit session expired. Start again with /edit ID.")
        return ConversationHandler.END
    if field == "name" and not 2 <= len(value) <= 100:
        await update.effective_message.reply_text("Enter a name between 2 and 100 characters.")
        return EDIT_VALUE
    if field == "phone" and not valid_phone(value):
        await update.effective_message.reply_text("Enter a valid phone number, e.g. +251912345678.")
        return EDIT_VALUE
    if field == "address" and not 2 <= len(value) <= 200:
        await update.effective_message.reply_text("Enter an address between 2 and 200 characters.")
        return EDIT_VALUE

    try:
        updated = edit_student(student_id, field, value)
    except PermissionError:
        await update.effective_message.reply_text("Close students.xlsx in Excel, then try again.")
        return ConversationHandler.END
    except Exception:
        logger.exception("Unable to update student")
        await update.effective_message.reply_text("The update failed. Please check the Excel file.")
        return ConversationHandler.END

    await update.effective_message.reply_text(
        f"Student {student_id} updated successfully." if updated else "Student ID not found."
    )
    context.user_data.pop("edit_student_id", None)
    context.user_data.pop("edit_field", None)
    return ConversationHandler.END


async def delete_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not await admin_only(update):
        return
    if not context.args or not context.args[0].isdigit():
        await update.effective_message.reply_text("Usage: /delete STUDENT_ID\nExample: /delete 1")
        return

    student_id = int(context.args[0])
    student = get_student(student_id)
    if not student:
        await update.effective_message.reply_text("Student ID not found.")
        return

    try:
        deleted = remove_student(student_id)
    except PermissionError:
        await update.effective_message.reply_text("Close students.xlsx in Excel, then try again.")
        return
    except Exception:
        logger.exception("Unable to delete student")
        await update.effective_message.reply_text("The deletion failed. Please check the Excel file.")
        return

    await update.effective_message.reply_text(
        f"Student {student_id} ({student['name']}) was deleted."
        if deleted else "Student ID not found."
    )


async def export_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not await admin_only(update):
        return
    backup = BASE_DIR / ("students_export_" + datetime.now().strftime("%Y%m%d_%H%M%S") + ".xlsx")
    try:
        ensure_excel()
        shutil.copy2(EXCEL_FILE, backup)
        with backup.open("rb") as file:
            await update.effective_message.reply_document(
                document=file,
                filename=backup.name,
                caption="Gete Registration Bot - Student Register",
            )
    except Exception:
        logger.exception("Unable to export Excel workbook")
        await update.effective_message.reply_text(
            "Export failed. Check that students.xlsx exists and is not locked."
        )
    finally:
        if backup.exists():
            backup.unlink()


# =========================================================
# ADMIN BROADCAST
# =========================================================

async def broadcast_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not await admin_only(update):
        return ConversationHandler.END

    students = read_students()
    recipients = sorted({
        s["chat_id"] for s in students if s.get("chat_id") is not None
    })
    if not recipients:
        await update.effective_message.reply_text(
            "There are no registered students with saved Telegram chat IDs yet. "
            "Students must register with this updated bot before they can receive broadcasts."
        )
        return ConversationHandler.END

    context.user_data.pop("broadcast_message", None)
    context.user_data["broadcast_recipient_count"] = len(recipients)
    await update.effective_message.reply_text(
        f"Broadcast preparation\n\n"
        f"Recipients with saved chat IDs: {len(recipients)}\n\n"
        "Send the announcement text you want to broadcast.\n"
        "It will not be sent until you confirm it.\n\n"
        "Use /cancel to stop."
    )
    return BROADCAST_MESSAGE


async def receive_broadcast_message(update: Update, context: ContextTypes.DEFAULT_TYPE):
    text = update.effective_message.text.strip()
    if not text:
        await update.effective_message.reply_text("The announcement cannot be empty. Send the text.")
        return BROADCAST_MESSAGE
    if len(text) > 4000:
        await update.effective_message.reply_text(
            "Please keep the announcement under 4,000 characters."
        )
        return BROADCAST_MESSAGE

    context.user_data["broadcast_message"] = text
    recipients = context.user_data.get("broadcast_recipient_count", 0)
    await update.effective_message.reply_text(
        "BROADCAST PREVIEW\n\n"
        f"{text}\n\n"
        f"Recipients: {recipients}\n\n"
        "Send YES to confirm and send it, or NO to cancel."
    )
    return BROADCAST_CONFIRM


async def confirm_broadcast(update: Update, context: ContextTypes.DEFAULT_TYPE):
    answer = update.effective_message.text.strip().casefold()
    if answer in {"no", "n", "cancel"}:
        context.user_data.pop("broadcast_message", None)
        context.user_data.pop("broadcast_recipient_count", None)
        await update.effective_message.reply_text("Broadcast cancelled. No messages were sent.")
        return ConversationHandler.END

    if answer not in {"yes", "y"}:
        await update.effective_message.reply_text("Reply YES to send or NO to cancel.")
        return BROADCAST_CONFIRM

    announcement = context.user_data.get("broadcast_message")
    if not announcement:
        await update.effective_message.reply_text("Broadcast expired. Start again with /broadcast.")
        return ConversationHandler.END

    # Re-read Excel at send time so recipients reflect the latest registrations.
    recipients = sorted({
        s["chat_id"] for s in read_students() if s.get("chat_id") is not None
    })
    if not recipients:
        await update.effective_message.reply_text("No saved student chat IDs were found.")
        return ConversationHandler.END

    await update.effective_message.reply_text(
        f"Sending announcement to {len(recipients)} registered chat(s)..."
    )

    sent = 0
    failed = 0
    bot = context.bot

    for chat_id in recipients:
        try:
            await bot.send_message(
                chat_id=chat_id,
                text="📢 Announcement from Gete Registration\n\n" + announcement,
            )
            sent += 1
            await asyncio.sleep(0.05)
        except RetryAfter as exc:
            try:
                delay = float(exc.retry_after) + 1
            except (TypeError, ValueError):
                delay = 2
            await asyncio.sleep(delay)
            try:
                await bot.send_message(
                    chat_id=chat_id,
                    text="📢 Announcement from Gete Registration\n\n" + announcement,
                )
                sent += 1
            except TelegramError:
                failed += 1
                logger.warning("Broadcast retry failed for chat ID %s", chat_id)
        except (Forbidden, TelegramError):
            failed += 1
            logger.warning("Broadcast failed for chat ID %s", chat_id)
        except Exception:
            failed += 1
            logger.exception("Unexpected broadcast failure for chat ID %s", chat_id)

    context.user_data.pop("broadcast_message", None)
    context.user_data.pop("broadcast_recipient_count", None)
    await update.effective_message.reply_text(
        "BROADCAST COMPLETE\n\n"
        f"Successful: {sent}\n"
        f"Failed: {failed}\n"
        f"Total recipients: {len(recipients)}\n\n"
        "Students who blocked the bot or have not started a private chat may not receive messages."
    )
    logger.info("Broadcast completed: %s sent, %s failed", sent, failed)
    return ConversationHandler.END


# =========================================================
# ERROR HANDLER / MAIN
# =========================================================

async def error_handler(update: object, context: ContextTypes.DEFAULT_TYPE):
    logger.error("An error occurred while processing an update", exc_info=context.error)


def main():
    ensure_excel()
    application = Application.builder().token(BOT_TOKEN).build()

    registration_handler = ConversationHandler(
        entry_points=[
            CommandHandler("start", start),
            CommandHandler("register", register),
            CommandHandler("edit", edit_command),
        ],
        states={
            NAME: [MessageHandler(filters.TEXT & ~filters.COMMAND, receive_name)],
            PHONE: [MessageHandler(filters.TEXT & ~filters.COMMAND, receive_phone)],
            ADDRESS: [MessageHandler(filters.TEXT & ~filters.COMMAND, receive_address)],
            REGISTRATION_CONFIRM: [MessageHandler(filters.TEXT & ~filters.COMMAND, confirm_registration)],
            EDIT_FIELD: [MessageHandler(filters.TEXT & ~filters.COMMAND, receive_edit_field)],
            EDIT_VALUE: [MessageHandler(filters.TEXT & ~filters.COMMAND, receive_edit_value)],
        },
        fallbacks=[CommandHandler("cancel", cancel)],
        allow_reentry=True,
    )

    broadcast_handler = ConversationHandler(
        entry_points=[CommandHandler("broadcast", broadcast_command)],
        states={
            BROADCAST_MESSAGE: [
                MessageHandler(filters.TEXT & ~filters.COMMAND, receive_broadcast_message)
            ],
            BROADCAST_CONFIRM: [
                MessageHandler(filters.TEXT & ~filters.COMMAND, confirm_broadcast)
            ],
        },
        fallbacks=[CommandHandler("cancel", cancel)],
        allow_reentry=True,
    )

    application.add_handler(registration_handler)
    application.add_handler(broadcast_handler)
    application.add_handler(CommandHandler("help", help_command))
    application.add_handler(CommandHandler("admin", admin_command))
    application.add_handler(CommandHandler("cancel", cancel))
    application.add_handler(CommandHandler("students", students_command))
    application.add_handler(CommandHandler("search", search_command))
    application.add_handler(CommandHandler("delete", delete_command))
    application.add_handler(CommandHandler("export", export_command))
    application.add_error_handler(error_handler)

    print("=" * 45)
    print("GETE REGISTRATION BOT")
    print(f"Excel: {EXCEL_FILE}")
    print(f"Configured admin IDs: {len(ADMIN_IDS)}")
    print("Status: Running")
    print("=" * 45)

    application.run_polling()


if __name__ == "__main__":
    main()
